from io import BytesIO

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf.errors import PdfReadError

from promptpilot_backend.document_service import (
    BasicDocumentParser,
    UrlIngestionService,
    _public_host,
)


def test_txt_and_csv_regression():
    parser = BasicDocumentParser()
    assert parser.parse("notes.txt", "text/plain", b"hello")[0].text == "hello"
    assert "A, B" in parser.parse("data.csv", "text/csv", b"A,B\n1,2")[0].text


def test_docx_paragraph_and_table():
    document = WordDocument()
    document.add_paragraph("Project purpose")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Feature"
    table.rows[0].cells[1].text = "Priority"
    stream = BytesIO()
    document.save(stream)
    parsed = BasicDocumentParser().parse(
        "requirements.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
    )
    assert any("Project purpose" in item.text for item in parsed)
    assert any("Feature" in item.text and "table 1" in item.provenance for item in parsed)


def test_xlsx_sheet_and_row_provenance():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Product", "Amount"])
    sheet.append(["Biryani", 10])
    stream = BytesIO()
    workbook.save(stream)
    parsed = BasicDocumentParser().parse(
        "sales.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        stream.getvalue(),
    )
    assert "Product=Biryani" in parsed[0].text
    assert 'sheet "Sales" — row 2' in parsed[0].provenance


def test_malformed_binary_fails_safely():
    with pytest.raises((PdfReadError, ValueError)):
        BasicDocumentParser().parse("bad.pdf", "application/pdf", b"not a pdf")


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://localhost", "http://127.0.0.1", "https://169.254.169.254"])
def test_url_ssrf_destinations_are_rejected(url):
    with pytest.raises(ValueError):
        UrlIngestionService().fetch(url)


def test_private_host_is_rejected():
    with pytest.raises(ValueError):
        _public_host("127.0.0.1")
