from __future__ import annotations

import csv
import hashlib
import io
import ipaddress
import re
import socket
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse
from uuid import UUID, uuid4

import httpx
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader
from sqlalchemy.orm import Session

from .config import get_settings
from .models import Document, DocumentChunk

ALLOWED_TYPES = {
    ".txt": "text/plain", ".csv": "text/csv", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAGIC = {".pdf": b"%PDF-", ".docx": b"PK", ".xlsx": b"PK"}


@dataclass(frozen=True)
class ParsedContent:
    text: str
    provenance: str


class DocumentParser:
    def parse(self, name: str, media_type: str, content: bytes) -> list[ParsedContent]:
        raise NotImplementedError


class BasicDocumentParser(DocumentParser):
    def parse(self, name: str, media_type: str, content: bytes) -> list[ParsedContent]:
        extension = Path(name).suffix.lower()
        if extension == ".txt":
            return [ParsedContent(content.decode("utf-8-sig", errors="strict"), f"{name} — text")]
        if extension == ".csv":
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="strict"))))
            return [ParsedContent("\n".join(", ".join(row) for row in rows), f"{name} — rows 1-{len(rows)}")]
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content), strict=True)
            return [ParsedContent(page.extract_text() or "", f"{name} — page {index}") for index, page in enumerate(reader.pages, 1)]
        if extension == ".docx":
            document = WordDocument(io.BytesIO(content))
            paragraphs = [ParsedContent(f"[Paragraph] {paragraph.text}", f"{name} — paragraph {index}") for index, paragraph in enumerate(document.paragraphs, 1) if paragraph.text.strip()]
            for table_index, table in enumerate(document.tables, 1):
                table_rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                if table_rows:
                    paragraphs.append(ParsedContent(f"[Table] {' | '.join(table_rows)}", f"{name} — table {table_index}"))
            return paragraphs
        if extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            result: list[ParsedContent] = []
            for sheet in workbook.worksheets:
                headers: list[Any] = list(next(sheet.iter_rows(values_only=True), ()))
                for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    values = [f"{headers[index] or f'Column {index + 1}'}={value}" for index, value in enumerate(row) if value is not None]
                    if values:
                        result.append(ParsedContent(f"[Sheet: {sheet.title}] [Row {row_index}] " + ", ".join(values), f'{name} — sheet "{sheet.title}" — row {row_index}'))
            return result
        raise ValueError(f"Parser for {extension} is not enabled")


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def chunk_text(parsed: list[ParsedContent], size: int = 1200, overlap: int = 120) -> list[tuple[str, str]]:
    output = []
    for item in parsed:
        text = normalize(item.text)
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            output.append((text[start:end], item.provenance))
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    return output


def _public_host(hostname: str) -> None:
    if hostname.lower().rstrip(".") in {"localhost", "metadata.google.internal", "metadata.google.com"}:
        raise ValueError("URL destination is not public")
    try:
        addresses = {info[4][0] for info in socket.getaddrinfo(hostname, None)}
    except socket.gaierror as error:
        raise ValueError("URL hostname could not be resolved") from error
    if any((address := ipaddress.ip_address(value)).is_private or address.is_loopback or address.is_link_local or address.is_reserved or address.is_unspecified for value in addresses):
        raise ValueError("URL destination is not public")


class UrlIngestionService:
    def __init__(self, document_service: DocumentService | None = None) -> None:
        self.document_service = document_service or DocumentService()

    def fetch(self, url: str) -> tuple[str, str, bytes]:
        current = url
        settings = get_settings()
        timeout = httpx.Timeout(settings.url_read_timeout, connect=settings.url_connect_timeout)
        for _ in range(settings.url_max_redirects + 1):
            parsed = urlparse(current)
            if parsed.scheme not in {"http", "https"} or not parsed.hostname:
                raise ValueError("Only public HTTP and HTTPS URLs are supported")
            _public_host(parsed.hostname)
            with httpx.Client(timeout=timeout, follow_redirects=False) as client:
                response = client.get(current, headers={"User-Agent": "PromptPilot/1.0"})
            if response.status_code in {301, 302, 303, 307, 308}:
                location = response.headers.get("location")
                if not location:
                    raise ValueError("Redirect response has no location")
                current = urljoin(current, location)
                continue
            if response.status_code >= 400:
                raise ValueError("URL could not be fetched")
            if len(response.content) > settings.url_max_response_bytes:
                raise ValueError("URL response exceeds configured size limit")
            return current, response.headers.get("content-type", "").split(";", 1)[0].lower(), response.content
        raise ValueError("URL redirect limit exceeded")

    def ingest_url(self, db: Session, project_id: UUID, url: str) -> Document:
        final_url, media_type, content = self.fetch(url)
        if media_type in {"text/html", "application/xhtml+xml"}:
            soup = BeautifulSoup(content, "html.parser")
            for node in soup(["script", "style", "nav", "footer", "header"]):
                node.decompose()
            title = soup.title.get_text(" ", strip=True) if soup.title else urlparse(final_url).netloc
            text = soup.get_text(" ", strip=True)
            return self.document_service.ingest_file(db, project_id, f"{title[:200]}.txt", "text/plain", text.encode(), source_type="url", source_url=final_url, provenance=f"{final_url} — extracted page text")
        extension = {value: key for key, value in ALLOWED_TYPES.items()}.get(media_type)
        if not extension:
            raise ValueError("URL content type is not supported")
        return self.document_service.ingest_file(db, project_id, f"download{extension}", media_type, content, source_type="url", source_url=final_url)


class DocumentService:
    def __init__(self, parser: DocumentParser | None = None) -> None:
        self.parser = parser or BasicDocumentParser()

    def ingest_file(self, db: Session, project_id: UUID, name: str, media_type: str, content: bytes, *, source_type: str = "file", source_url: str | None = None, provenance: str | None = None) -> Document:
        extension = Path(name).suffix.lower()
        if extension not in ALLOWED_TYPES or media_type not in {ALLOWED_TYPES[extension], "application/octet-stream"}:
            raise ValueError("Unsupported or mismatched file type")
        if len(content) > get_settings().max_upload_bytes:
            raise ValueError("File exceeds configured size limit")
        if extension in MAGIC and not content.startswith(MAGIC[extension]):
            raise ValueError("File signature does not match its type")
        checksum = hashlib.sha256(content).hexdigest()
        storage_key = f"documents/{project_id}/{uuid4().hex}{extension}"
        storage_root = Path(get_settings().storage_path)
        (storage_root / storage_key).parent.mkdir(parents=True, exist_ok=True)
        (storage_root / storage_key).write_bytes(content)
        document = Document(project_id=project_id, name=Path(name).name, media_type=media_type, source_type=source_type, source_url=source_url, storage_key=storage_key, size_bytes=len(content), checksum=checksum, status="processing")
        db.add(document)
        db.flush()
        try:
            for index, (text, item_provenance) in enumerate(chunk_text(self.parser.parse(name, media_type, content))):
                db.add(DocumentChunk(document_id=document.id, chunk_index=index, content=text, provenance=provenance or item_provenance, character_count=len(text), processing_version="document-processing-v2"))
            document.status = "processed"
        except Exception as error:
            document.status = "failed"
            document.error_message = str(error)[:500]
        db.commit()
        db.refresh(document)
        return document
